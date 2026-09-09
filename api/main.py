import os,json,secrets,shutil,re,hashlib
from pathlib import Path
from datetime import datetime
from fastapi import FastAPI,Request,HTTPException,UploadFile,File
from fastapi.responses import FileResponse,JSONResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
BASE=Path(__file__).resolve().parent.parent
DATA=Path(os.getenv('ERP_DATA_DIR',BASE/'runtime-data')).resolve(); DEFAULT=BASE/'data-default'; DATA.mkdir(parents=True,exist_ok=True)
for p in DEFAULT.rglob('*'):
 q=DATA/p.relative_to(DEFAULT)
 if p.is_dir(): q.mkdir(parents=True,exist_ok=True)
 elif not q.exists(): q.parent.mkdir(parents=True,exist_ok=True); shutil.copy2(p,q)
DB=DATA/'data.json'; CAT=DATA/'catalog.json'; CFG=json.loads((DATA/'template_sources.json').read_text(encoding='utf8'))
DEFAULT_ROLES={'ADMIN': ['DASHBOARD', 'ITEM', 'VENDOR', 'PO', 'APPROVE', 'RECEIVE', 'IQC', 'INVENTORY', 'NCR', 'AUDIT', 'COST', 'USER_ADMIN', 'SYSTEM_SETTINGS', 'MASTER', 'EXPORT', 'REPORTS'], 'MANAGER': ['DASHBOARD', 'VENDOR', 'PO', 'APPROVE', 'RECEIVE', 'IQC', 'INVENTORY', 'NCR', 'AUDIT', 'COST', 'EXPORT', 'REPORTS'], 'PURCHASING': ['DASHBOARD', 'ITEM', 'VENDOR', 'PO', 'COST', 'MASTER', 'EXPORT', 'REPORTS'], 'WAREHOUSE': ['DASHBOARD', 'RECEIVE', 'INVENTORY'], 'QC': ['DASHBOARD', 'IQC', 'NCR']}
def role_rights(d,role): return d.get('roles',DEFAULT_ROLES).get(role,[])
SESS={}; app=FastAPI(title='Evolution ERP V11 FastAPI')
def hash_password(p):
 salt=os.urandom(16); return 'pbkdf2_sha256$260000$'+salt.hex()+'$'+hashlib.pbkdf2_hmac('sha256',p.encode(),salt,260000).hex()
def verify_password(p,h):
 try:
  _,n,salt,digest=h.split('$'); return secrets.compare_digest(hashlib.pbkdf2_hmac('sha256',p.encode(),bytes.fromhex(salt),int(n)).hex(),digest)
 except: return False
def migrate_data(d):
 emails={'ADM-001':'admin@evolution.ca','PUR-001':'purchasing@evolution.ca','WH-001':'warehouse@evolution.ca','QC-001':'quality@evolution.ca','MGR-001':'manager@evolution.ca'}; changed=False
 for u in d.get('users',[]):
  if not u.get('email'): u['email']=emails.get(u.get('employeeId'),u.get('employeeId','user').lower()+'@evolution.ca'); changed=True
  if 'password' in u: u['passwordHash']=hash_password(u.pop('password')); changed=True
  for k,v in [('preferredLanguage','vi'),('themePreference','light'),('lastLoginAt','')]:
   if k not in u:u[k]=v;changed=True
 if 'roles' not in d:d['roles']=DEFAULT_ROLES;changed=True
 defaults={'systemTitle':'Evolution Technologies ERP','companyName':'Evolution Technologies Vietnam','defaultLanguage':'vi','defaultTheme':'light','timezone':'Asia/Ho_Chi_Minh','logoFile':'company_logo.png'}
 for k,v in defaults.items():
  if k not in d['settings']:d['settings'][k]=v;changed=True
 return changed
def read():
 d=json.loads(DB.read_text(encoding='utf8'))
 if migrate_data(d): write(d)
 return d
def write(d):
 t=DB.with_suffix('.tmp'); t.write_text(json.dumps(d,ensure_ascii=False,indent=2),encoding='utf8'); t.replace(DB)
def user(req,p=None):
 token=req.headers.get('authorization','').removeprefix('Bearer '); u=SESS.get(token)
 if not u: raise HTTPException(401,'Please log in')
 if p and p not in role_rights(read(),u['role']): raise HTTPException(403,'Permission denied')
 return u
def audit(d,u,m,k,a,details=''): d['audit'].insert(0,{'time':datetime.now().astimezone().isoformat(),'userName':u['name'],'employeeId':u['employeeId'],'role':u['role'],'module':m,'referenceNo':k,'action':a,'details':details})
def doc(prefix,n): return f"{prefix}-{datetime.now():%Y%m%d}-{n:04d}"
def next_po(d):
 suf=f'{datetime.now():%m%Y}'; nums=[int(m.group(1)) for p in d['purchaseOrders'] if (m:=re.match(rf'^PO-(\d+)-{suf}$',p['poNumber']))]; return f"PO-{max(nums,default=0)+1:02d}-{suf}"
def stats(d,p,l):
 rec=sum(r['deliveredQty'] for r in d['receivings'] if r['poNumber']==p['poNumber'] and r['lineNo']==l['lineNo']); q=l['totalQty']; return {**l,'receivedQty':rec,'remainingQty':max(0,q-rec),'lineStatus':'OPEN' if rec==0 else 'FULLY_RECEIVED' if rec>=q else 'PARTIALLY_RECEIVED'}
def enriched(d):
 out=[]
 for p in d['purchaseOrders']:
  ls=[stats(d,p,l) for l in p['lines']]; rec=sum(x['receivedQty'] for x in ls); tot=sum(x['totalQty'] for x in ls)
  out.append({**p,'lines':ls,'receivingStatus':'OPEN' if rec==0 else 'FULLY_RECEIVED' if rec>=tot else 'PARTIALLY_RECEIVED'})
 return out
@app.exception_handler(HTTPException)
async def he(req,e): return JSONResponse({'error':e.detail},e.status_code)
@app.post('/api/login')
async def login(req:Request):
 x=await req.json(); d=read(); email=str(x.get('email','')).strip().lower(); u=next((v for v in d['users'] if v.get('email','').lower()==email and v.get('active')),None)
 if not u or not verify_password(str(x.get('password','')),u.get('passwordHash','')):
  d['audit'].insert(0,{'time':datetime.now().astimezone().isoformat(),'userName':email,'employeeId':'','role':'','module':'AUTH','referenceNo':email,'action':'LOGIN_FAILED','details':req.client.host if req.client else ''}); write(d); raise HTTPException(401,'Invalid email or password')
 u['lastLoginAt']=datetime.now().isoformat(); safe={k:u.get(k,'') for k in ['name','employeeId','email','role','preferredLanguage','themePreference']}; t=secrets.token_hex(24); SESS[t]=safe; audit(d,safe,'AUTH',email,'LOGIN',req.client.host if req.client else ''); write(d); return {'token':t,'user':safe,'permissions':role_rights(d,safe['role']),'settings':d['settings']}
@app.get('/api/me')
def me(req:Request): u=user(req); d=read(); return {'user':u,'permissions':role_rights(d,u['role']),'settings':d['settings']}
@app.post('/api/logout')
def logout(req:Request): SESS.pop(req.headers.get('authorization','').removeprefix('Bearer '),None); return {'ok':True}
@app.get('/api/data')
def data(req:Request):
 user(req); d=read(); c=json.loads(CAT.read_text(encoding='utf8')); return {**d,'users':[],'vendors':[v for v in d['vendors'] if not v.get('templateCode')]+c['vendors'],'aluminumPartMaster':c['items'],'poTemplates':CFG['templates'],'purchaseOrders':enriched(d),'nextPONumber':next_po(d)}
@app.post('/api/vendor')
async def vendor(req:Request):
 u=user(req,'VENDOR'); x=await req.json(); d=read(); name=x.get('vendorName','').strip().upper()
 if not name: raise HTTPException(400,'Vendor name required')
 vid='VND-'+re.sub('[^A-Z0-9]+','-',name)
 if any(v['vendorName']==name for v in d['vendors']): raise HTTPException(409,'Vendor exists')
 v={'vendorId':vid,'vendorName':name,'email':x.get('email',''),'active':True}; d['vendors'].append(v); audit(d,u,'Vendor',vid,'CREATE'); write(d); return v
@app.post('/api/catalog/reload')
def reload_catalog(req:Request):
 user(req,'MASTER'); wb=load_workbook(DATA/'EXTRUSION_INFORMATION.xlsx',data_only=True); vs=wb['VENDOR INFORMATION']; ms=wb['Master list']; vendors=[]; items=[]
 for r in range(4,vs.max_row+1):
  code=str(vs[f'B{r}'].value or '').strip()
  if code: vendors.append({'vendorId':'EXT-'+code,'vendorName':code,'description':str(vs[f'C{r}'].value or ''),'attn':str(vs[f'D{r}'].value or ''),'email':str(vs[f'E{r}'].value or '').replace('mailto:',''),'mobile':str(vs[f'F{r}'].value or ''),'paymentTerms':'30 Days after delivery','deliveryTerms':'DDP - Evolution factory','active':True,'templateCode':'ALUMINUM_EXTRUSION'})
 old=json.loads(CAT.read_text(encoding='utf8')); images={i['fabricationItemCode']:i.get('imageFile','') for i in old['items']}
 for r in range(4,ms.max_row+1):
  drawing=str(ms[f'B{r}'].value or '').strip(); code=str(ms[f'E{r}'].value or '').strip()
  if drawing and code: items.append({'drawingName':drawing,'drawingCode':drawing,'extrusionItemCode':str(ms[f'C{r}'].value or ''),'material':str(ms[f'D{r}'].value or ''),'fabricationItemCode':code,'partName':str(ms[f'F{r}'].value or ''),'model':str(ms[f'H{r}'].value or ''),'cuttingLength':str(ms[f'I{r}'].value or ''),'barLength':float(ms[f'J{r}'].value or 0),'pcsPerBar':float(ms[f'K{r}'].value or 0),'cuttingPrice':float(ms[f'L{r}'].value or 0),'fabricationPrice':float(ms[f'M{r}'].value or 0),'cleaningPrice':float(ms[f'N{r}'].value or 0),'active':True,'templateCode':'ALUMINUM_EXTRUSION','imageFile':images.get(code,'')})
 CAT.write_text(json.dumps({'vendors':vendors,'items':items,'loadedAt':datetime.now().isoformat()},ensure_ascii=False,indent=2),encoding='utf8'); return {'vendors':len(vendors),'items':len(items)}
@app.get('/api/price-history')
def price(req:Request,itemCode:str=''):
 user(req); d=read(); out=[{'poNumber':p['poNumber'],'date':p['orderDate'],'vendorId':p['vendorId'],'unitPrice':l['totalUnitPrice']} for p in d['purchaseOrders'] for l in p['lines'] if l.get('itemCode')==itemCode]; return sorted(out,key=lambda x:x['date'],reverse=True)
@app.post('/api/po')
async def create_po(req:Request):
 u=user(req,'PO'); x=await req.json(); d=read(); cat=json.loads(CAT.read_text(encoding='utf8'))
 if not x.get('lines'): raise HTTPException(400,'Add at least one PO item')
 lines=[]
 for i,l in enumerate(x['lines'],1):
  if x['poType']=='NON_INVENTORY':
   qty=float(l.get('qty',0)); pr=float(l.get('unitPrice',0)); lines.append({'lineNo':i,**l,'partName':l.get('description',''),'itemCode':l.get('itemCode',''),'qtyBar':qty,'totalQty':qty,'pcsPerBar':1,'cuttingPrice':0,'fabricationPrice':0,'cleaningPrice':0,'totalUnitPrice':pr,'amount':qty*pr,'lineStatus':'OPEN'})
  else:
   m=next((v for v in cat['items'] if v['drawingCode']==l.get('drawingCode') and v['fabricationItemCode']==l.get('itemCode')),None)
   if not m: raise HTTPException(400,'Selected item is not in Master list')
   bars=float(l.get('qtyBar',0)); total=bars*float(m['pcsPerBar']); cut=float(l.get('cuttingPrice',0)); fab=float(l.get('fabricationPrice',0)); clean=float(l.get('cleaningPrice',0)); unitp=cut+fab+clean
   lines.append({'lineNo':i,**m,**{k:l.get(k) for k in ['onHandAtOrder','dailyRequirement','workingDays','targetCoverageDays','monthlyDemand','coverageDays','suggestedOrderQty','stockStatus']},'qtyBar':bars,'totalQty':total,'cuttingPrice':cut,'fabricationPrice':fab,'cleaningPrice':clean,'totalUnitPrice':unitp,'amount':total*unitp,'itemCode':m['fabricationItemCode'],'lineStatus':'OPEN'})
 sub=sum(l['amount'] for l in lines); vat=float(d['settings']['vatPercent']); po={'poNumber':next_po(d),'poType':x['poType'],'templateCode':x['templateCode'],'vendorId':x['vendorId'],'orderDate':datetime.now().date().isoformat(),'expectedDate':x.get('expectedDate',''),'currency':d['settings']['currency'],'paymentTerms':x.get('paymentTerms',''),'deliveryTerms':x.get('deliveryTerms',''),'department':x.get('department',''),'status':'DRAFT','approvalStatus':'DRAFT','approverEmail':d['settings']['approverEmail'],'createdBy':u['name'],'createdById':u['employeeId'],'createdAt':datetime.now().isoformat(),'lines':lines,'subtotal':sub,'vatPercent':vat,'vatAmount':sub*vat/100,'grandTotal':sub*(1+vat/100)}
 d['purchaseOrders'].insert(0,po); audit(d,u,'PO',po['poNumber'],'CREATE',f'{len(lines)} lines'); write(d); return po
@app.post('/api/po/{no}/submit')
def submit(no:str,req:Request):
 u=user(req,'PO'); d=read(); p=next((p for p in d['purchaseOrders'] if p['poNumber']==no),None)
 if not p or p['status']!='DRAFT': raise HTTPException(409,'Only Draft PO can be submitted')
 p['status']=p['approvalStatus']='PENDING_APPROVAL'; p['submittedAt']=datetime.now().isoformat(); audit(d,u,'PO',no,'SUBMIT_FOR_APPROVAL',p['approverEmail']); write(d); return {'ok':True,'note':'Use Manager/Admin Approvals screen'}
def export_po(p,d):
 inv=p['poType']=='INVENTORY'; tpl=DATA/('EXTRUSION_INFORMATION.xlsx' if inv else 'NON_INVENTORY_FORM.xlsx'); wb=load_workbook(tpl); ws=wb.worksheets[0]; wb._sheets=[ws]; ws.title=re.sub(r'[\\/*?:\[\]]','_',p['poNumber'])[:31]; c=json.loads(CAT.read_text(encoding='utf8')); v=next((x for x in c['vendors']+d['vendors'] if x['vendorId']==p['vendorId']),{})
 if inv:
  for cell,val in {'B5':v.get('description',v.get('vendorName','')),'B6':v.get('attn',''),'B7':v.get('email',''),'B8':v.get('mobile',''),'R5':p['poNumber'],'Q6':'Date: '+p['orderDate'],'Q7':'Delivery Terms: '+p.get('deliveryTerms',''),'Q8':'Payment Terms: '+p.get('paymentTerms',''),'A9':'Currency : VND'}.items(): ws[cell]=val
  for i,l in enumerate(p['lines'][:7],13):
   vals=[i-12,f"{l.get('drawingCode','')}\n({l.get('itemCode','')})",l.get('extrusionItemCode',''),l.get('material',''),l.get('itemCode',''),l.get('partName',''),'IMAGE',l.get('model',''),l.get('cuttingLength',''),l.get('qtyBar',0),l.get('barLength',0),l.get('pcsPerBar',0),l.get('totalQty',0),l.get('cuttingPrice',0),l.get('fabricationPrice',0),l.get('cleaningPrice',0),l.get('totalUnitPrice',0),l.get('amount',0)]
   cols=['A','B','C','D','E','F','G','H','I','J','K','L','M','N','O','P','Q','R']
   for col,val in zip(cols,vals): ws[f'{col}{i}']=val
   im=DATA/'uploads'/l.get('imageFile','');
   if im.exists(): img=XLImage(im); img.width=90; img.height=55; ws.add_image(img,f'G{i}')
  ws['R20']=p['subtotal']; ws['R21']=p['vatAmount']; ws['R22']=p['grandTotal']
 else:
  for cell,val in {'B6':v.get('description',v.get('vendorName','')),'B7':v.get('attn',''),'B8':v.get('email',''),'B9':v.get('mobile',''),'G6':p['poNumber'],'G7':p['orderDate'],'H8':p.get('deliveryTerms',''),'H9':p.get('paymentTerms',''),'A10':'Currency : VND'}.items(): ws[cell]=val
  for i,l in enumerate(p['lines'][:4],13):
   for col,val in zip(['A','B','D','E','F','G','H','I'],[i-12,l.get('description',''),l.get('model',''),l.get('specification',''),l.get('unit','PCS'),l.get('totalQty',0),l.get('totalUnitPrice',0),l.get('amount',0)]): ws[f'{col}{i}']=val
  ws['I17']=p['subtotal']; ws['I18']=p['vatAmount']; ws['I19']=p['grandTotal']
 out=DATA/'archive'/f"{p['poNumber']}.xlsx"; out.parent.mkdir(exist_ok=True); wb.save(out); return out
@app.post('/api/po/{no}/approve')
def approve(no:str,req:Request):
 u=user(req,'APPROVE'); d=read(); p=next((p for p in d['purchaseOrders'] if p['poNumber']==no),None)
 if not p or p['status']!='PENDING_APPROVAL': raise HTTPException(409,'PO is not pending approval')
 p['status']=p['approvalStatus']='APPROVED'; p['approvedBy']=u['name']; p['approvedById']=u['employeeId']; p['approvedAt']=datetime.now().isoformat(); export_po(p,d); audit(d,u,'PO',no,'APPROVE'); write(d); return {'ok':True,'fileName':no+'.xlsx'}
@app.post('/api/po/{no}/reject')
async def reject(no:str,req:Request):
 u=user(req,'APPROVE'); x=await req.json(); d=read(); p=next((p for p in d['purchaseOrders'] if p['poNumber']==no),None)
 if not p or p['status']!='PENDING_APPROVAL': raise HTTPException(409,'PO is not pending approval')
 p['status']=p['approvalStatus']='REJECTED'; p['rejectedBy']=u['name']; p['rejectionReason']=x.get('reason',''); audit(d,u,'PO',no,'REJECT',p['rejectionReason']); write(d); return {'ok':True}
@app.get('/api/po/{no}/export')
def exp(no:str,req:Request):
 user(req,'EXPORT'); d=read(); p=next((p for p in d['purchaseOrders'] if p['poNumber']==no),None)
 if not p: raise HTTPException(404,'PO not found')
 out=DATA/'archive'/f'{no}.xlsx';
 if not out.exists(): export_po(p,d)
 return FileResponse(out,filename=out.name)
@app.post('/api/receive')
async def receive(req:Request):
 u=user(req,'RECEIVE'); x=await req.json(); d=read(); p=next((p for p in d['purchaseOrders'] if p['poNumber']==x['poNumber']),None); l=next((l for l in p['lines'] if l['lineNo']==int(x['lineNo'])),None) if p else None
 if not p or p['status']!='APPROVED' or not l: raise HTTPException(409,'Only Approved PO can be received')
 st=stats(d,p,l); qty=float(x['deliveredQty']);
 if qty<=0 or qty>st['remainingQty']: raise HTTPException(409,'Invalid receiving quantity')
 no=doc('RCV',d['counters']['receiving']); d['counters']['receiving']+=1; d['receivings'].insert(0,{'receivingNo':no,'poNumber':p['poNumber'],'lineNo':l['lineNo'],'itemCode':l['itemCode'],'partName':l['partName'],'deliveredQty':qty,'lotNumber':str(x.get('lotNumber','')).upper(),'deliveryNoteNo':x.get('deliveryNoteNo',''),'status':'IQC_PENDING' if p['poType']=='INVENTORY' else 'DEPARTMENT_CONFIRMATION','receivedBy':u['name'],'createdAt':datetime.now().isoformat()}); audit(d,u,'Receiving',no,'CREATE'); write(d); return {'receivingNo':no}
@app.post('/api/iqc/{receiving_no}')
async def iqc(receiving_no:str,req:Request):
 u=user(req,'IQC'); x=await req.json(); d=read(); r=next((v for v in d['receivings'] if v['receivingNo']==receiving_no),None); a=float(x['acceptedQty']); j=float(x['rejectedQty'])
 if not r or r['status']!='IQC_PENDING' or a+j!=r['deliveredQty']: raise HTTPException(400,'Invalid IQC result')
 r['status']='PASS' if j==0 else 'FAIL' if a==0 else 'PARTIAL_ACCEPT'; no=doc('IQC',d['counters']['iqc']); d['counters']['iqc']+=1; d['iqc'].insert(0,{'iqcNo':no,'receivingNo':r['receivingNo'],'poNumber':r['poNumber'],'lineNo':r['lineNo'],'inspectedQty':float(x.get('inspectedQty',r['deliveredQty'])),'acceptedQty':a,'rejectedQty':j,'result':r['status'],'defectCode':x.get('defectCode',''),'severity':x.get('severity',''),'disposition':x.get('disposition',''),'remarks':x.get('remarks',''),'checkedBy':u['name'],'checkedAt':datetime.now().isoformat()})
 if a>0:
  b=next((v for v in d['inventory'] if v['itemCode']==r['itemCode'] and v['lotNumber']==r['lotNumber']),None)
  if not b: b={'itemCode':r['itemCode'],'partName':r['partName'],'lotNumber':r['lotNumber'],'onHandQty':0}; d['inventory'].append(b)
  b['onHandQty']+=a
 if j>0: n=doc('NCR',d['counters']['ncr']); d['counters']['ncr']+=1; d['ncr'].insert(0,{'ncrNo':n,'iqcNo':no,'poNumber':r['poNumber'],'lineNo':r['lineNo'],'itemCode':r['itemCode'],'rejectedQty':j,'status':'OPEN'})
 audit(d,u,'IQC',no,r['status']); write(d); return {'result':r['status']}

def admin_user(req): return user(req,'USER_ADMIN')
@app.get('/api/admin/users')
def list_users(req:Request):
 admin_user(req); return [{k:u.get(k,'') for k in ['name','employeeId','email','role','active','preferredLanguage','themePreference','lastLoginAt']} for u in read()['users']]
@app.post('/api/admin/users')
async def add_user(req:Request):
 actor=admin_user(req); x=await req.json(); d=read(); email=x.get('email','').strip().lower()
 if not email or '@' not in email: raise HTTPException(400,'Valid email required')
 if any(u.get('email','').lower()==email for u in d['users']): raise HTTPException(409,'Email already exists')
 role=x.get('role','PURCHASING')
 if role not in d['roles']: raise HTTPException(400,'Invalid role')
 u={'name':x.get('name','').strip() or email.split('@')[0],'employeeId':x.get('employeeId','').strip().upper() or 'USR-'+secrets.token_hex(3).upper(),'email':email,'role':role,'active':True,'passwordHash':hash_password(x.get('password') or 'ChangeMe123!'),'preferredLanguage':x.get('preferredLanguage','vi'),'themePreference':x.get('themePreference','light'),'lastLoginAt':''};d['users'].append(u);audit(d,actor,'ADMIN',email,'CREATE_USER');write(d);return {'ok':True}
@app.put('/api/admin/users/{email}')
async def edit_user(email:str,req:Request):
 actor=admin_user(req);x=await req.json();d=read();u=next((v for v in d['users'] if v.get('email','').lower()==email.lower()),None)
 if not u:raise HTTPException(404,'User not found')
 for k in ['name','role','active','preferredLanguage','themePreference']:
  if k in x:u[k]=x[k]
 if x.get('password'):u['passwordHash']=hash_password(x['password'])
 audit(d,actor,'ADMIN',email,'UPDATE_USER');write(d);return {'ok':True}
@app.get('/api/admin/roles')
def list_roles(req:Request): admin_user(req);return read()['roles']
@app.put('/api/admin/roles')
async def edit_roles(req:Request):
 actor=admin_user(req);x=await req.json();d=read();d['roles']=x;audit(d,actor,'ADMIN','ROLES','UPDATE_PERMISSIONS');write(d);return {'ok':True}
@app.get('/api/admin/audit')
def admin_audit(req:Request):admin_user(req);return read()['audit'][:500]
@app.get('/api/audit')
def audit_view(req:Request): user(req,'AUDIT'); return read()['audit'][:1000]
@app.post('/api/admin/roles/{role_name}')
def add_role(role_name:str,req:Request):
 actor=admin_user(req);d=read();name=role_name.strip().upper().replace(' ','_')
 if not name:raise HTTPException(400,'Role name required')
 if name in d['roles']:raise HTTPException(409,'Role already exists')
 d['roles'][name]=[];audit(d,actor,'ADMIN',name,'CREATE_ROLE');write(d);return {'ok':True,'role':name}
@app.delete('/api/admin/roles/{role_name}')
def remove_role(role_name:str,req:Request):
 actor=admin_user(req);d=read();name=role_name.upper()
 if name=='ADMIN':raise HTTPException(409,'ADMIN role cannot be removed')
 if any(u.get('role')==name for u in d['users']):raise HTTPException(409,'Role is assigned to a user')
 if name not in d['roles']:raise HTTPException(404,'Role not found')
 del d['roles'][name];audit(d,actor,'ADMIN',name,'DELETE_ROLE');write(d);return {'ok':True}
@app.get('/api/admin/settings')
def admin_settings(req:Request):admin_user(req);return read()['settings']
@app.put('/api/admin/settings')
async def edit_settings(req:Request):
 actor=admin_user(req);x=await req.json();d=read();d['settings'].update(x);audit(d,actor,'ADMIN','SETTINGS','UPDATE_SETTINGS');write(d);return d['settings']
@app.put('/api/profile')
async def profile(req:Request):
 current=user(req);x=await req.json();d=read();u=next(v for v in d['users'] if v.get('email')==current.get('email'))
 for k in ['preferredLanguage','themePreference']:
  if k in x:u[k]=x[k];current[k]=x[k]
 write(d);return {'ok':True,'user':current}

@app.get('/health')
def health(): return {'status':'ok','engine':'FastAPI','nodejs':False}
app.mount('/',StaticFiles(directory=BASE/'public',html=True),name='public')
